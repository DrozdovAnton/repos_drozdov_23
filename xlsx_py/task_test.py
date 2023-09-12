import openpyxl as xl
import pathlib
import os
import time

from datetime import datetime
from docxtpl import DocxTemplate
from docxcompose.composer import Composer
from docx import Document


# FROM XLSX TO PYTHON LISTS
ls = [] # общий лист данных из xlsx

ls_86 = []
ls_1 = []
ls_2 = []
ls_3 = []


# LIST FOR DIR
all_dir = [] # все доступные папки в директории

if os.path.isfile('M-CHAT-R.xlsx'):
    n_st = int(input('Введите номер строки Excel, с которой надо начать автозаполнение: '))
    n_end = int(input('Введите номер строки Excel, до которой надо заполнять формы: '))

    table = xl.load_workbook('M-CHAT-R.xlsx')
    sheet = table.active

    for i in range(n_st, n_end+1): # from xlsx to list
        current_ls = []
        for j in range(10): # кол-во столбцов, необходимых для извлечения
            if j != 7:
                if j == 2:
                    current_ls.append(str(int(sheet[i][j].value)) + '/20')
                elif j == 0:
                    dt = sheet[i][j].value
                    current_ls.append(str(dt.strftime('%d.%m.%Y %H:%M:%S')))
                elif j == 4:
                    dt = sheet[i][j].value
                    current_ls.append(str(dt.strftime('%d.%m.%Y')))
                else:
                    current_ls.append(str(sheet[i][j].value))
            else:
                continue
        ls.append(current_ls)

    for i in ls: # sort from list to lists
        if i[7] == 'ДГП 86':
            ls_86.append(i)
        elif i[7] == 'ДГП 86 филиал 1':
            ls_1.append(i)
        elif i[7] == 'ДГП 86 филиал 2':
            ls_2.append(i)
        elif i[7] == 'ДГП 86 филиал 3':
            ls_3.append(i)

    def list_to_docx(ls): # from lists to docx
        doc = DocxTemplate('шаблон.docx')
        for l in ls:
            context = { 'childrennm' : l[3], 'bday' : l[4],
                        'parentnm' : l[5], 'phone' : l[6], 'email' : l[1],
                        'numberorg' : l[7], 'docnm' : l[8], 'points' : l[2], 
                        'time' : l[0]}
            doc.render(context)

            frm_time = l[0].replace('.', '-').replace(':', '-')

            if l[7] == 'ДГП 86':
                str_path = './86/' + l[3] + ' ' + frm_time + '.docx'
                doc.save(str_path)

            elif l[7] == 'ДГП 86 филиал 1':
                str_path = './86_1/' + l[3]  + ' ' + frm_time + '.docx'
                doc.save(str_path)

            elif l[7] == 'ДГП 86 филиал 2':
                str_path = './86_2/' + l[3] + ' ' + frm_time + '.docx'
                doc.save(str_path)

            elif l[7] == 'ДГП 86 филиал 3':
                str_path = './86_3/' + l[3] + ' ' + frm_time + '.docx'
                doc.save(str_path)

    def file_to_files(dir_name): # from docxs to docx

        empty = Document()
        composer = Composer(empty)

        ls_files = os.listdir(dir_name)

        for i in ls_files:
            path = os.path.dirname(os.path.abspath(i)) + chr(92) + dir_name + chr(92) + i
            composer.append(Document(path))
        
        name = dir_name + '_merge.docx'
        composer.save(name)

    def final_build():
        for dir in os.listdir():
            if '.' not in dir:
                all_dir.append(dir)
        
        for dir in all_dir:
            file_to_files(dir)

    list_to_docx(ls_86)
    list_to_docx(ls_1)
    list_to_docx(ls_2)
    list_to_docx(ls_3)

    final_build()

else:
    print('Таблица в директории отсутствует!', 'Необходимо внести в директорию актуальную таблицу!', sep='\n')
    time.sleep(6.0)